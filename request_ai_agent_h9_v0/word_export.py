"""DOCX export for the browser-rendered analysis request preview.

This module preserves the existing Preview DOM -> /api/export/word contract and
replaces the previous hand-written OpenXML renderer with python-docx.

Runtime dependency:
    python-docx>=1.1,<2
"""

from __future__ import annotations

from io import BytesIO
import re
from typing import Any, Mapping, Sequence

try:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError as exc:  # Keep the application importable until export is used.
    Document = None  # type: ignore[assignment]
    _PYTHON_DOCX_IMPORT_ERROR: ImportError | None = exc
else:
    _PYTHON_DOCX_IMPORT_ERROR = None


DOCX_MIMETYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
DEFAULT_WORD_FILENAME = "analysis_request.docx"

_FONT_NAME = "Malgun Gothic"
_TEXT_COLOR = "202124"
_MUTED_COLOR = "5F6368"
_ACCENT_COLOR = "3F4B59"
_BORDER_COLOR = "D7DBDF"
_HEADER_FILL = "E9EDF1"
_LABEL_FILL = "F4F5F6"
_ALT_ROW_FILL = "FAFBFC"
_SECTION_RULE_COLOR = "AEB5BC"
_PORTRAIT_CONTENT_WIDTH_CM = 17.0
_TABLE_BULLET_GLYPHS = "•●▪◦∙‣⁃◉○◎◆◇▶►▸‧・·"
_LEADING_TABLE_BULLET_RE = re.compile(rf"(?m)^[ \t]*[{_TABLE_BULLET_GLYPHS}]+[ \t]*")
_SPACED_TABLE_BULLET_RE = re.compile(rf"[ \t]+[{_TABLE_BULLET_GLYPHS}]+[ \t]+")

# Long/free-text fields read better as a full-width row than as a 2-up key/value grid.
_NARRATIVE_LABELS = {
    "요청 내용",
    "의사결정 활용",
    "추가 결과 요청",
    "기준 제품 대비 형상 차이",
    "해석을 요청하게 된 배경",
    "해석으로 확인하고 싶은 내용",
}


def word_filename() -> str:
    """Use the existing stable filename contract."""
    return DEFAULT_WORD_FILENAME


def _require_python_docx() -> None:
    if _PYTHON_DOCX_IMPORT_ERROR is not None or Document is None:
        raise RuntimeError(
            "Word export requires the 'python-docx' package. "
            "Install it with: pip install 'python-docx>=1.1,<2'"
        ) from _PYTHON_DOCX_IMPORT_ERROR


def _text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).replace("\x00", "").strip()


def _table_text(value: Any) -> str:
    """Remove presentation-only bullet glyphs without changing stored values."""

    text = _LEADING_TABLE_BULLET_RE.sub("", _text(value))
    return _SPACED_TABLE_BULLET_RE.sub(" ", text)


def _items(value: Any) -> list[Any]:
    return value if isinstance(value, list) else []


def _mapping(value: Any) -> Mapping[str, Any]:
    return value if isinstance(value, Mapping) else {}


def _set_east_asia_font(run: Any, name: str = _FONT_NAME) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def _set_run_character_spacing(run: Any, spacing_twips: int) -> None:
    """Apply compact character spacing; Word stores this in twentieths of a point."""

    r_pr = run._element.get_or_add_rPr()
    spacing = r_pr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        r_pr.append(spacing)
    spacing.set(qn("w:val"), str(int(spacing_twips)))


def _set_run_style(
    run: Any,
    *,
    size_pt: float = 9.5,
    bold: bool = False,
    color: str = _TEXT_COLOR,
    character_spacing_twips: int | None = None,
) -> None:
    _set_east_asia_font(run)
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    if character_spacing_twips is not None:
        _set_run_character_spacing(run, character_spacing_twips)


def _set_paragraph_spacing(
    paragraph: Any,
    *,
    before: float = 0,
    after: float = 0,
    line: float = 1.08,
) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def _set_cell_shading(cell: Any, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_margins(cell: Any, *, top: int = 90, start: int = 110, bottom: int = 90, end: int = 110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_fixed_width(table: Any, total_width_cm: float) -> None:
    """Keep tables inside the portrait content width instead of expanding the page."""

    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    width_twips = round(float(total_width_cm) / 2.54 * 1440)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(width_twips))


def _set_table_borders(table: Any, *, color: str = _BORDER_COLOR, size: str = "5") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "start", "bottom", "end", "insideH", "insideV"):
        tag = f"w:{edge}"
        border = borders.find(qn(tag))
        if border is None:
            border = OxmlElement(tag)
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), size)
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), color)


def _set_repeat_table_header(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def _set_paragraph_bottom_border(paragraph: Any, *, color: str = _SECTION_RULE_COLOR, size: str = "6") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "5")
    bottom.set(qn("w:color"), color)


def _set_cell_text(
    cell: Any,
    value: Any,
    *,
    bold: bool = False,
    color: str = _TEXT_COLOR,
    size_pt: float = 9.0,
    align: Any = None,
    compact: bool = False,
    character_spacing_twips: int | None = None,
) -> None:
    if align is None:
        align = WD_ALIGN_PARAGRAPH.LEFT
    text = _table_text(value)
    if not text:
        text = "-"
    paragraph = cell.paragraphs[0]
    paragraph.clear()
    paragraph.style = "Normal"
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is not None:
        p_pr.remove(num_pr)
    paragraph.alignment = align
    _set_paragraph_spacing(paragraph, before=0, after=0, line=0.95 if compact else 1.05)
    run = paragraph.add_run(text)
    _set_run_style(
        run,
        size_pt=size_pt,
        bold=bold,
        color=color,
        character_spacing_twips=character_spacing_twips,
    )
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if compact:
        _set_cell_margins(cell, top=45, start=35, bottom=45, end=35)
    else:
        _set_cell_margins(cell)


def _configure_document(document: Any) -> None:
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.8)

    normal = document.styles["Normal"]
    normal.font.name = _FONT_NAME
    normal.font.size = Pt(10.0)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), _FONT_NAME)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.08

    document.core_properties.title = "해석 의뢰서"

    # A restrained footer gives multi-page requests a stable document cue.
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_paragraph_spacing(paragraph)
    run = paragraph.add_run("해석 의뢰서  |  ")
    _set_run_style(run, size_pt=8.0, color=_MUTED_COLOR)
    fld_simple = OxmlElement("w:fldSimple")
    fld_simple.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld_simple)


def _add_document_title(document: Any) -> None:
    paragraph = document.add_paragraph()
    _set_paragraph_spacing(paragraph, before=0, after=13, line=1.0)
    run = paragraph.add_run("해석 의뢰서")
    _set_run_style(run, size_pt=21.0, bold=True, color=_TEXT_COLOR)
    _set_paragraph_bottom_border(paragraph, color=_ACCENT_COLOR, size="12")


def _request_identity_value(value: Any, *, request_no: bool = False) -> str:
    text = _text(value)
    if request_no:
        for prefix in ("해석 의뢰 번호:", "해석의뢰 번호:"):
            if text.startswith(prefix):
                text = text[len(prefix):].strip()
                break
    if text == "해석 의뢰를 시작해 주세요.":
        return "-"
    return text or "-"


def _add_request_identity(document: Any, request_title: Any, request_no: Any) -> None:
    table = document.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    _set_table_fixed_width(table, _PORTRAIT_CONTENT_WIDTH_CM)
    _set_table_borders(table, color=_BORDER_COLOR, size="5")
    widths = (Cm(3.1), Cm(13.9))
    rows = (
        ("의뢰 제목", _request_identity_value(request_title)),
        ("의뢰 번호", _request_identity_value(request_no, request_no=True)),
    )
    for row, values in zip(table.rows, rows):
        for cell, width in zip(row.cells, widths):
            cell.width = width
        _set_cell_shading(row.cells[0], _LABEL_FILL)
        _set_cell_text(row.cells[0], values[0], bold=True, color=_MUTED_COLOR, size_pt=9.0)
        _set_cell_text(row.cells[1], values[1], size_pt=9.0)

    spacer = document.add_paragraph()
    _set_paragraph_spacing(spacer, before=0, after=7, line=1.0)


def _add_section_heading(document: Any, title: str, index: int) -> None:
    paragraph = document.add_paragraph()
    _set_paragraph_spacing(paragraph, before=17 if index > 1 else 2, after=9, line=1.0)
    number_run = paragraph.add_run(f"{index:02d}  ")
    _set_run_style(number_run, size_pt=10.0, bold=True, color=_ACCENT_COLOR)
    title_run = paragraph.add_run(title or "구분")
    _set_run_style(title_run, size_pt=13.0, bold=True, color=_TEXT_COLOR)
    _set_paragraph_bottom_border(paragraph, color=_SECTION_RULE_COLOR, size="6")


def _add_group_heading(document: Any, title: Any) -> None:
    paragraph = document.add_paragraph()
    _set_paragraph_spacing(paragraph, before=9, after=5, line=1.0)
    run = paragraph.add_run(_text(title) or "구분")
    _set_run_style(run, size_pt=10.5, bold=True, color=_ACCENT_COLOR)


def _add_stale_notice(document: Any, text: Any) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = True
    _set_table_borders(table, color="C6C6C6", size="5")
    cell = table.cell(0, 0)
    _set_cell_shading(cell, "F7F7F7")
    _set_cell_text(cell, text, bold=True, color=_MUTED_COLOR, size_pt=9.0)


def _add_key_value_table(
    document: Any,
    fields: Sequence[tuple[str, str]],
    *,
    page_width_cm: float = _PORTRAIT_CONTENT_WIDTH_CM,
) -> None:
    """Render concise business data as two label/value pairs per row."""

    if not fields:
        return

    table = document.add_table(rows=0, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    _set_table_fixed_width(table, page_width_cm)
    _set_table_borders(table, color=_BORDER_COLOR, size="5")

    label_width = 3.1
    value_width = max(1.0, (page_width_cm - label_width * 2) / 2)
    widths = (Cm(label_width), Cm(value_width), Cm(label_width), Cm(value_width))

    for offset in range(0, len(fields), 2):
        row = table.add_row()
        for cell, width in zip(row.cells, widths):
            cell.width = width

        left_label, left_value = fields[offset]
        _set_cell_shading(row.cells[0], _LABEL_FILL)
        _set_cell_text(row.cells[0], left_label, bold=True, color=_MUTED_COLOR, size_pt=9.0)
        _set_cell_text(row.cells[1], left_value, size_pt=9.0)

        if offset + 1 < len(fields):
            right_label, right_value = fields[offset + 1]
            _set_cell_shading(row.cells[2], _LABEL_FILL)
            _set_cell_text(row.cells[2], right_label, bold=True, color=_MUTED_COLOR, size_pt=9.0)
            _set_cell_text(row.cells[3], right_value, size_pt=9.0)
        else:
            value_cell = row.cells[1].merge(row.cells[2]).merge(row.cells[3])
            value_cell.width = Cm(page_width_cm - label_width)
            _set_cell_text(value_cell, left_value, size_pt=9.0)

def _add_narrative_field(document: Any, label: str, value: str) -> None:
    """Give long request context enough horizontal and vertical reading space."""

    table = document.add_table(rows=2, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    _set_table_fixed_width(table, _PORTRAIT_CONTENT_WIDTH_CM)
    _set_table_borders(table, color=_BORDER_COLOR, size="5")
    for row in table.rows:
        row.cells[0].width = Cm(_PORTRAIT_CONTENT_WIDTH_CM)

    label_cell = table.cell(0, 0)
    _set_cell_shading(label_cell, _LABEL_FILL)
    _set_cell_text(label_cell, label, bold=True, color=_ACCENT_COLOR, size_pt=9.0)

    value_cell = table.cell(1, 0)
    value_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    _set_cell_margins(value_cell, top=135, start=145, bottom=150, end=145)
    paragraph = value_cell.paragraphs[0]
    paragraph.clear()
    paragraph.style = "Normal"
    _set_paragraph_spacing(paragraph, line=1.05)
    _set_run_style(paragraph.add_run(value or "-"), size_pt=9.0)

def _add_narrative_fields(document: Any, fields: Sequence[tuple[str, str]]) -> None:
    for index, (label, value) in enumerate(fields):
        if index:
            spacer = document.add_paragraph()
            _set_paragraph_spacing(spacer, after=3, line=1.0)
        _add_narrative_field(document, label, value)


def _table_width_weights(
    headers: Sequence[str],
    rows: Sequence[Sequence[str]],
    *,
    compact: bool = False,
) -> list[float]:
    if not headers:
        return []
    weights: list[float] = []
    for index, header in enumerate(headers):
        lengths = [len(header or "-")]
        for row in rows[:20]:
            if index < len(row):
                lengths.append(len(row[index] or "-"))
        max_len = max(lengths)
        minimum_weight = 4 if compact else 7
        maximum_weight = 18 if compact else 26
        weights.append(float(min(maximum_weight, max(minimum_weight, max_len))))
    return weights


def _column_widths_cm(
    headers: Sequence[str],
    rows: Sequence[Sequence[str]],
    total_width_cm: float,
    *,
    compact: bool = False,
) -> list[float]:
    weights = _table_width_weights(headers, rows, compact=compact)
    if not weights:
        return []
    total_weight = sum(weights) or 1.0
    raw = [total_width_cm * weight / total_weight for weight in weights]
    # Avoid unusably narrow columns while preserving the total width approximately.
    minimum = (0.95 if len(headers) >= 7 else 1.35) if compact else (1.8 if len(headers) >= 7 else 2.2)
    clamped = [max(minimum, width) for width in raw]
    scale = total_width_cm / sum(clamped)
    return [width * scale for width in clamped]


def _preview_column_widths_cm(table_key: str, headers: Sequence[str], total_width_cm: float) -> list[float]:
    """Mirror fixed SCREEN-06 preview column ratios where the DOM defines them."""

    column_count = len(headers)
    if table_key == "case_matrix":
        case_ratios = {
            "No.": 5,
            "No": 5,
            "형상": 17,
            "운전 조건": 16,
            "열교환기 사양": 24,
            "공간 환경 조건": 16,
            "취출 공기 조건": 16,
        }
        ratios = tuple(case_ratios.get(header, 16) for header in headers)
        ratio_total = sum(ratios)
        return [total_width_cm * ratio / ratio_total for ratio in ratios]

    ratios = {
        "geometry": (14, 26, 60),
        "operating_conditions": (20, 18, 62),
        "heat_exchanger_conditions": (14, 22, 22, 22, 10, 10),
    }.get(table_key)
    if ratios is None or len(ratios) != column_count:
        return []
    ratio_total = sum(ratios)
    return [total_width_cm * ratio / ratio_total for ratio in ratios]


def _add_data_table(
    document: Any,
    table_block: Mapping[str, Any],
    *,
    page_width_cm: float = 17.0,
    section_title: str = "",
) -> None:
    caption = _text(table_block.get("caption"))
    table_key = _text(table_block.get("key"))
    headers = [_text(item) for item in _items(table_block.get("headers"))]
    if headers and not headers[0]:
        headers[0] = {
            "operating_conditions": "운전 구분",
            "heat_exchanger_conditions": "사양 구분",
        }.get(table_key, "구분")
    headers = [item or "-" for item in headers]
    rows = [[_text(item) or "-" for item in _items(row)] for row in _items(table_block.get("rows"))]

    if caption:
        paragraph = document.add_paragraph()
        _set_paragraph_spacing(paragraph, before=2, after=5, line=1.0)
        run = paragraph.add_run(caption)
        _set_run_style(run, size_pt=9.5, bold=True, color=_MUTED_COLOR)

    column_count = max(len(headers), max((len(row) for row in rows), default=0))
    if column_count <= 0:
        return

    is_case_matrix = (
        table_key == "case_matrix"
        or "case matrix" in caption.lower()
        or "case matrix" in _text(section_title).lower()
    )

    table = document.add_table(rows=0, cols=column_count)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    _set_table_fixed_width(table, page_width_cm)
    _set_table_borders(table)

    normalized_headers = headers + ["-"] * (column_count - len(headers))
    normalized_rows = [row + ["-"] * (column_count - len(row)) for row in rows]
    widths = _preview_column_widths_cm(table_key, normalized_headers, page_width_cm)
    if not widths:
        widths = _column_widths_cm(normalized_headers, normalized_rows, page_width_cm, compact=is_case_matrix)

    if normalized_headers:
        header_row = table.add_row()
        _set_repeat_table_header(header_row)
        for index, value in enumerate(normalized_headers):
            cell = header_row.cells[index]
            if index < len(widths):
                cell.width = Cm(widths[index])
            _set_cell_shading(cell, _HEADER_FILL)
            _set_cell_text(
                cell,
                value,
                bold=True,
                color=_TEXT_COLOR,
                size_pt=7.5 if is_case_matrix else 9.0,
                compact=is_case_matrix,
                character_spacing_twips=-3 if is_case_matrix else None,
            )

    for row_index, values in enumerate(normalized_rows):
        row = table.add_row()
        for index, value in enumerate(values):
            cell = row.cells[index]
            if index < len(widths):
                cell.width = Cm(widths[index])
            if row_index % 2:
                _set_cell_shading(cell, _ALT_ROW_FILL)
            _set_cell_text(
                cell,
                value,
                size_pt=7.0 if is_case_matrix else 9.0,
                align=WD_ALIGN_PARAGRAPH.CENTER if is_case_matrix and index == 0 else None,
                compact=is_case_matrix,
                character_spacing_twips=-3 if is_case_matrix else None,
            )
            if is_case_matrix:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

def _section_blocks(section: Mapping[str, Any]) -> list[Mapping[str, Any]]:
    return [_mapping(block) for block in _items(section.get("blocks"))]


def _field_pairs(section: Mapping[str, Any]) -> list[tuple[str, str]]:
    return [
        (_text(block.get("label")) or "항목", _text(block.get("value")) or "-")
        for block in _section_blocks(section)
        if _text(block.get("type")) == "field"
    ]


def _find_section_index(sections: Sequence[Mapping[str, Any]], title: str) -> int | None:
    for index, section in enumerate(sections):
        if _text(section.get("title")) == title:
            return index
    return None


def _find_table(section: Mapping[str, Any], key: str) -> Mapping[str, Any] | None:
    for block in _section_blocks(section):
        if _text(block.get("type")) == "table" and _text(block.get("key")) == key:
            return block
    return None


def _render_blocks(
    document: Any,
    blocks: Sequence[Mapping[str, Any]],
    *,
    section_title: str,
    page_width_cm: float = _PORTRAIT_CONTENT_WIDTH_CM,
) -> None:
    """Render a section in source order while keeping fields under their group."""

    field_buffer: list[tuple[str, str]] = []

    def flush_fields() -> None:
        nonlocal field_buffer
        if field_buffer:
            _add_key_value_table(document, field_buffer, page_width_cm=page_width_cm)
            field_buffer = []

    for block in blocks:
        kind = _text(block.get("type"))
        if kind == "field":
            field_buffer.append((_text(block.get("label")) or "항목", _text(block.get("value")) or "-"))
            continue

        flush_fields()
        if kind == "group":
            _add_group_heading(document, block.get("title"))
        elif kind == "stale":
            _add_stale_notice(document, _text(block.get("text")) or "-")
        elif kind == "table":
            _add_data_table(document, block, page_width_cm=page_width_cm, section_title=section_title)

    flush_fields()


def _add_case_summary(document: Any, row_count: int) -> None:
    paragraph = document.add_paragraph()
    _set_paragraph_spacing(paragraph, before=0, after=7, line=1.05)
    run = paragraph.add_run(f"총 {row_count}개 Case로 구성됩니다.")
    _set_run_style(run, size_pt=9.0, color=_MUTED_COLOR)


def _case_table_for_word(case_table: Mapping[str, Any]) -> dict[str, Any]:
    """Keep the SCREEN-06 matrix shape while removing requested generic names."""

    headers = [_text(item) for item in _items(case_table.get("headers"))]
    rows = [[_text(item) for item in _items(row)] for row in _items(case_table.get("rows"))]
    generic_prefixes = {"운전 조건": "운전", "열교환기 사양": "사양"}

    for header, prefix in generic_prefixes.items():
        if header not in headers:
            continue
        column_index = headers.index(header)
        pattern = re.compile(rf"^{re.escape(prefix)}\s*\d+\s*(?:[{_TABLE_BULLET_GLYPHS}]+\s*)?")
        for row in rows:
            if column_index >= len(row):
                continue
            without_name = pattern.sub("", row[column_index]).strip()
            if header == "열교환기 사양":
                without_name = re.sub(r"^(F&T|MC)(?=\S)", r"\1\n", without_name)
            row[column_index] = without_name or "-"

    return {
        "type": "table",
        "key": "case_matrix",
        "caption": _text(case_table.get("caption")),
        "headers": headers,
        "rows": rows,
    }


def _render_semantic_document(document: Any, sections: Sequence[Mapping[str, Any]]) -> None:
    """Recompose SCREEN-06 data into a document-first engineering narrative."""

    requester_index = _find_section_index(sections, "의뢰자 정보")
    overview_index = _find_section_index(sections, "요청 내용")
    geometry_index = _find_section_index(sections, "해석 제품")
    conditions_index = _find_section_index(sections, "해석 조건")
    case_index = _find_section_index(sections, "Case Matrix")
    consumed: set[int] = set()
    section_number = 0

    requester_fields = _field_pairs(sections[requester_index]) if requester_index is not None else []
    overview_fields = _field_pairs(sections[overview_index]) if overview_index is not None else []
    narrative_fields = [item for item in overview_fields if item[0] in _NARRATIVE_LABELS]
    request_fields = [item for item in overview_fields if item[0] not in _NARRATIVE_LABELS]

    if requester_index is not None:
        consumed.add(requester_index)
    if overview_index is not None:
        consumed.add(overview_index)

    if requester_fields or request_fields:
        section_number += 1
        _add_section_heading(document, "의뢰 개요", section_number)
        if requester_fields:
            _add_group_heading(document, "의뢰자 정보")
            _add_key_value_table(document, requester_fields)
        if request_fields:
            _add_group_heading(document, "의뢰 기본 정보 및 해석유형")
            _add_key_value_table(document, request_fields)

    if narrative_fields:
        section_number += 1
        _add_section_heading(document, "해석 요청사항", section_number)
        _add_narrative_fields(document, narrative_fields)

    if geometry_index is not None:
        consumed.add(geometry_index)
        section_number += 1
        _add_section_heading(document, "해석 대상", section_number)
        _render_blocks(document, _section_blocks(sections[geometry_index]), section_title="해석 대상")

    if conditions_index is not None:
        consumed.add(conditions_index)
        section_number += 1
        _add_section_heading(document, "해석 조건", section_number)
        _render_blocks(document, _section_blocks(sections[conditions_index]), section_title="해석 조건")

    # Preserve additive or future SCREEN-06 sections without forcing them into a
    # known semantic group. Case Matrix remains last as the request conclusion.
    for index, section in enumerate(sections):
        if index in consumed or index == case_index:
            continue
        title = _text(section.get("title")) or "추가 정보"
        blocks = _section_blocks(section)
        if not blocks:
            continue
        section_number += 1
        _add_section_heading(document, title, section_number)
        _render_blocks(document, blocks, section_title=title)

    if case_index is None:
        return

    consumed.add(case_index)
    case_section = sections[case_index]
    case_table = _find_table(case_section, "case_matrix")
    case_blocks = _section_blocks(case_section)

    section_number += 1
    _add_section_heading(document, "Case 구성", section_number)
    if case_table is not None:
        _add_case_summary(document, len(_items(case_table.get("rows"))))
        _add_data_table(
            document,
            _case_table_for_word(case_table),
            page_width_cm=_PORTRAIT_CONTENT_WIDTH_CM,
            section_title="Case Matrix",
        )

    remaining_case_blocks = [
        block
        for block in case_blocks
        if not (_text(block.get("type")) == "table" and _text(block.get("key")) == "case_matrix")
    ]
    if remaining_case_blocks:
        _render_blocks(document, remaining_case_blocks, section_title="Case Matrix")



def build_word_docx(preview: Mapping[str, Any]) -> bytes:
    """Create a document-first DOCX from the serialized current preview DOM.

    Input contract is the browser Preview plus additive request identity:
        {"request_title": str, "request_no": str, "sections": [...]}

    The Preview/API contract and supported block types remain unchanged. The
    presentation is intentionally recomposed for an engineering request reader.
    """
    _require_python_docx()

    document = Document()
    _configure_document(document)
    _add_document_title(document)
    metadata = _mapping(preview.get("metadata"))
    _add_request_identity(
        document,
        preview.get("request_title", metadata.get("request_title")),
        preview.get("request_no", metadata.get("request_no")),
    )

    sections = [
        _mapping(section)
        for section in _items(preview.get("sections"))
        if _text(_mapping(section).get("title")) or _items(_mapping(section).get("blocks"))
    ]
    _render_semantic_document(document, sections)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
