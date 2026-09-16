from io import BytesIO

import pytest
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn

from request_ai_agent_h9_v0.ui import HTML_TEMPLATE
from request_ai_agent_h9_v0.word_export import build_word_docx


def _preview_payload():
    return {
        "request_title": "HA / RAC / Sample / 유동해석",
        "request_no": "REQ-001",
        "sections": [
            {
                "title": "의뢰자 정보",
                "blocks": [
                    {"type": "field", "label": "사업부", "value": "H&A"},
                    {"type": "field", "label": "부서", "value": "RAC개발실"},
                    {"type": "field", "label": "요청자", "value": "홍길동"},
                    {"type": "field", "label": "직급", "value": "책임연구원"},
                ],
            },
            {
                "title": "요청 내용",
                "blocks": [
                    {"type": "field", "label": "의뢰 유형", "value": "신규"},
                    {"type": "field", "label": "프로젝트명(PMS)", "value": "PMS-01"},
                    {"type": "field", "label": "개발 등급", "value": "A"},
                    {"type": "field", "label": "NPI 단계", "value": "DV"},
                    {"type": "field", "label": "모델명(Model Suffix)", "value": "MODEL-A"},
                    {"type": "field", "label": "희망 완료일", "value": "2026-10-30"},
                    {"type": "field", "label": "해석유형", "value": "유동해석"},
                    {"type": "field", "label": "해석을 요청하게 된 배경", "value": "배경 설명"},
                    {"type": "field", "label": "해석으로 확인하고 싶은 내용", "value": "확인 내용"},
                ],
            },
            {
                "title": "해석 제품",
                "blocks": [
                    {
                        "type": "table",
                        "key": "geometry",
                        "headers": ["구분", "총조립도 도면번호 (NPDM MCAD)", "Base 대비 변경점"],
                        "rows": [["Base", "DRAW-A", "기존 형상"], ["비교 1", "DRAW-B", "토출부 변경"]],
                    }
                ],
            },
            {
                "title": "해석 조건",
                "blocks": [
                    {"type": "group", "title": "운전 조건"},
                    {
                        "type": "table",
                        "key": "operating_conditions",
                        "headers": ["", "팬 개수", "팬 회전 설정"],
                        "rows": [["운전 1", "2", "모든 팬 동일 · 900 RPM"]],
                    },
                    {"type": "group", "title": "열교환기 사양"},
                    {
                        "type": "table",
                        "key": "heat_exchanger_conditions",
                        "headers": ["", "HEX Type", "관 직경(Pi) / 채널 폭(Width)", "Fin type", "열 수", "FPI / FPDM"],
                        "rows": [["사양 1", "Fin-tube", "7 mm", "Louver", "2", "18"]],
                    },
                    {"type": "group", "title": "공간 환경 조건"},
                    {"type": "field", "label": "실내 온도", "value": "25 °C"},
                    {"type": "field", "label": "실내 상대습도", "value": "50 %"},
                    {"type": "group", "title": "취출 공기 조건"},
                    {"type": "field", "label": "취출 온도", "value": "10 °C"},
                    {"type": "field", "label": "취출 상대습도", "value": "70 %"},
                ],
            },
            {
                "title": "Case Matrix",
                "blocks": [
                    {
                        "type": "table",
                        "key": "case_matrix",
                        "headers": ["No.", "형상", "운전 조건", "열교환기 사양", "공간 환경 조건", "취출 공기 조건"],
                        "rows": [["1", "Base", "운전 1", "사양 1", "25 °C / 50 %", "10 °C / 70 %"]],
                    }
                ],
            },
        ],
    }


def _all_table_text(document):
    return "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)


def test_preview_serializer_preserves_condition_headings_and_table_identity():
    serializer = HTML_TEMPLATE.split("function previewDomForWord()", 1)[1].split(
        "async function exportWordFromPreview", 1
    )[0]

    for title in ("운전 조건", "열교환기 사양", "공간 환경 조건", "취출 공기 조건"):
        assert f"<h5 data-preview-group-title>{title}</h5>" in HTML_TEMPLATE
    assert "[data-preview-group-title]" in serializer
    assert '.filter(node => !node.matches(\'[data-preview-label]\') || !node.closest(\'table\'))' in serializer
    assert 'return {type:"group", title:clean(node.textContent)}' in serializer
    assert 'key:clean(node.dataset.previewTable)' in serializer


def test_word_recomposes_screen_six_as_a_reader_focused_engineering_request():
    document = Document(BytesIO(build_word_docx(_preview_payload())))
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text]

    assert paragraphs == [
        "해석 의뢰서",
        "01  의뢰 개요",
        "의뢰자 정보",
        "의뢰 기본 정보 및 해석유형",
        "02  해석 요청사항",
        "03  해석 대상",
        "04  해석 조건",
        "운전 조건",
        "열교환기 사양",
        "공간 환경 조건",
        "취출 공기 조건",
        "05  Case 구성",
        "총 1개 Case로 구성됩니다.",
    ]

    all_text = _all_table_text(document) + "\n" + "\n".join(paragraphs)
    for expected in (
        "H&A",
        "홍길동",
        "PMS-01",
        "유동해석",
        "배경 설명",
        "확인 내용",
        "DRAW-B",
        "토출부 변경",
        "모든 팬 동일 · 900 RPM",
        "Louver",
        "25 °C",
        "70 %",
        "Case 구성",
    ):
        assert expected in all_text


def test_word_uses_neutral_business_tables_and_keeps_screen_six_column_ratios():
    document = Document(BytesIO(build_word_docx(_preview_payload())))

    identity_table = document.tables[0]
    requester_table = document.tables[1]
    operating_table = next(table for table in document.tables if table.cell(0, 1).text == "팬 개수")
    specification_table = next(table for table in document.tables if table.cell(0, 1).text == "HEX Type")

    for table in (identity_table, requester_table, operating_table):
        borders = table._tbl.tblPr.find(qn("w:tblBorders"))
        assert borders is not None
        assert {border.get(qn("w:val")) for border in borders} == {"single"}

    identity_label_shading = identity_table.rows[0].cells[0]._tc.tcPr.find(qn("w:shd"))
    header_shading = operating_table.rows[0].cells[0]._tc.tcPr.find(qn("w:shd"))
    assert identity_label_shading is not None
    assert identity_label_shading.get(qn("w:fill")) == "F4F5F6"
    assert header_shading is not None
    assert header_shading.get(qn("w:fill")) == "E9EDF1"

    operating_widths = [cell.width.cm for cell in operating_table.rows[0].cells]
    specification_widths = [cell.width.cm for cell in specification_table.rows[0].cells]
    assert operating_widths == pytest.approx([3.4, 3.06, 10.54], abs=0.02)
    assert specification_widths == pytest.approx([2.38, 3.74, 3.74, 3.74, 1.7, 1.7], abs=0.02)

    narrative_tables = [table for table in document.tables if len(table.columns) == 1]
    assert [[row.cells[0].text for row in table.rows] for table in narrative_tables] == [
        ["해석을 요청하게 된 배경", "배경 설명"],
        ["해석으로 확인하고 싶은 내용", "확인 내용"],
    ]


def test_wide_case_matrix_moves_to_a_readable_landscape_page():
    document = Document(BytesIO(build_word_docx(_preview_payload())))

    assert len(document.sections) == 2
    assert document.sections[0].orientation == WD_ORIENT.PORTRAIT
    assert document.sections[0].page_width.cm == pytest.approx(21.0, abs=0.02)
    assert document.sections[1].orientation == WD_ORIENT.LANDSCAPE
    assert document.sections[1].page_width.cm == pytest.approx(29.7, abs=0.02)
    assert document.sections[1].page_height.cm == pytest.approx(21.0, abs=0.02)

    case_table = document.tables[-1]
    assert case_table.rows[0].cells[0].text == "No."
    assert case_table.rows[1].cells[1].text == "Base"
    assert sum(cell.width.cm for cell in case_table.rows[0].cells) == pytest.approx(25.7, abs=0.1)
