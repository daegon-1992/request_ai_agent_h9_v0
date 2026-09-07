"""DOCX export for the browser-rendered analysis request preview.

This module preserves the existing Preview DOM -> /api/export/word contract and
replaces the previous hand-written OpenXML renderer with python-docx.

Runtime dependency:
    python-docx>=1.1,<2
"""

from __future__ import annotations

from io import BytesIO
from typing import Any, Mapping, Sequence

try:
    from docx import Document
    from docx.enum.section import WD_ORIENT, WD_SECTION
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError as exc:  # Keep the application importable until export is used.
    Document = None  # type: ignore[assignment]
    _PYTHON_DOCX_IMPORT_ERROR: ImportError | None = exc
else:
    _PYTHON_DOCX_IMPORT_ERROR = None


DOCX_MIMETYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
DEFAULT_WORD_FILENAME = "analysis_request.docx"

_FONT_NAME = "Malgun Gothic"
_TEXT_COLOR = "222222"
_MUTED_COLOR = "6B6B6B"
_BORDER_COLOR = "D7D7D7"
_LABEL_FILL = "F4F4F4"
_HEADER_FILL = "EEEEEE"
_SECTION_RULE_COLOR = "BDBDBD"

# Long/free-text fields read better as a full-width row than as a 2-up key/value grid.
_NARRATIVE_LABELS = {
    "요청 내용",
    "의사결정 활용",
    "추가 결과 요청",
    "기준 제품 대비 형상 차이",
}


def word_filename() -> str:
    """Use the existing stable filename contract."""
    return DEFAULT_WORD_FILENAME


def _require_python_docx() -> None:
    if _PYTHON_DOCX_IMPORT_ERROR is not None or Document is None:
        raise RuntimeError(
            "Word export requires the 'python-docx' package. "
            "Install it with: pip install 'python-docx>=1.1,<2'"
        ) from _PYTHON_DOCX_IMPORT_ERROR


def _text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).replace("\x00", "").strip()


def _items(value: Any) -> list[Any]:
    return value if isinstance(value, list) else []


def _mapping(value: Any) -> Mapping[str, Any]:
    return value if isinstance(value, Mapping) else {}


def _set_east_asia_font(run: Any, name: str = _FONT_NAME) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def _set_run_style(
    run: Any,
    *,
    size_pt: float = 9.5,
    bold: bool = False,
    color: str = _TEXT_COLOR,
) -> None:
    _set_east_asia_font(run)
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def _set_paragraph_spacing(
    paragraph: Any,
    *,
    before: float = 0,
    after: float = 0,
    line: float = 1.08,
) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def _set_cell_shading(cell: Any, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_margins(cell: Any, *, top: int = 90, start: int = 110, bottom: int = 90, end: int = 110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_cell_border(cell: Any, *, color: str = _BORDER_COLOR, size: str = "5") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "start", "bottom", "end", "insideH", "insideV"):
        tag = f"w:{edge}"
        border = borders.find(qn(tag))
        if border is None:
            border = OxmlElement(tag)
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), size)
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), color)


def _set_table_borders(table: Any, *, color: str = _BORDER_COLOR, size: str = "5") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "start", "bottom", "end", "insideH", "insideV"):
        tag = f"w:{edge}"
        border = borders.find(qn(tag))
        if border is None:
            border = OxmlElement(tag)
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), size)
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), color)


def _set_repeat_table_header(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def _prevent_row_split(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def _set_paragraph_bottom_border(paragraph: Any, *, color: str = _SECTION_RULE_COLOR, size: str = "6") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "5")
    bottom.set(qn("w:color"), color)


def _set_cell_text(
    cell: Any,
    value: Any,
    *,
    bold: bool = False,
    color: str = _TEXT_COLOR,
    size_pt: float = 9.0,
    align: Any = None,
    allow_blank: bool = False,
) -> None:
    if align is None:
        align = WD_ALIGN_PARAGRAPH.LEFT
    text = _text(value)
    if not text and not allow_blank:
        text = "-"
    paragraph = cell.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = align
    _set_paragraph_spacing(paragraph, before=0, after=0, line=1.05)
    run = paragraph.add_run(text)
    _set_run_style(run, size_pt=size_pt, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _set_cell_margins(cell)


def _configure_document(document: Any) -> None:
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.8)

    normal = document.styles["Normal"]
    normal.font.name = _FONT_NAME
    normal.font.size = Pt(9.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), _FONT_NAME)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.08

    # A restrained footer keeps the output document-like without adding data not
    # present in the request state.
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_paragraph_spacing(paragraph)
    run = paragraph.add_run("해석 의뢰서  ·  ")
    _set_run_style(run, size_pt=8.0, color=_MUTED_COLOR)
    fld_simple = OxmlElement("w:fldSimple")
    fld_simple.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld_simple)


def _add_document_title(document: Any) -> None:
    paragraph = document.add_paragraph()
    _set_paragraph_spacing(paragraph, before=0, after=16, line=1.0)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run("해석 의뢰서")
    _set_run_style(run, size_pt=20.0, bold=True, color=_TEXT_COLOR)
    _set_paragraph_bottom_border(paragraph, color="8A8A8A", size="10")


def _add_section_heading(document: Any, title: str, index: int) -> None:
    paragraph = document.add_paragraph()
    _set_paragraph_spacing(paragraph, before=10 if index > 1 else 0, after=7, line=1.0)
    paragraph.paragraph_format.keep_with_next = True
    number_run = paragraph.add_run(f"{index:02d}  ")
    _set_run_style(number_run, size_pt=10.5, bold=True, color=_MUTED_COLOR)
    title_run = paragraph.add_run(title or "구분")
    _set_run_style(title_run, size_pt=11.5, bold=True, color=_TEXT_COLOR)
    _set_paragraph_bottom_border(paragraph)


def _add_stale_notice(document: Any, text: Any) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = True
    _set_table_borders(table, color="C6C6C6", size="5")
    cell = table.cell(0, 0)
    _set_cell_shading(cell, "F7F7F7")
    _set_cell_text(cell, text, bold=True, color=_MUTED_COLOR, size_pt=9.0)


def _is_narrative_field(label: str, value: str) -> bool:
    return label in _NARRATIVE_LABELS or "\n" in value or len(value) >= 70


def _add_field_grid(document: Any, fields: Sequence[tuple[str, str]]) -> None:
    if not fields:
        return

    table = document.add_table(rows=0, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    _set_table_borders(table)

    compact_pending: tuple[str, str] | None = None

    def add_compact_row(left: tuple[str, str], right: tuple[str, str] | None = None) -> None:
        row = table.add_row()
        _prevent_row_split(row)
        left_label, left_value = left
        right_label, right_value = right if right is not None else ("", "")

        cells = row.cells
        widths = (Cm(2.8), Cm(5.7), Cm(2.8), Cm(5.7))
        for cell, width in zip(cells, widths):
            cell.width = width

        _set_cell_shading(cells[0], _LABEL_FILL)
        _set_cell_text(cells[0], left_label, bold=True, color=_MUTED_COLOR, size_pt=8.8)
        _set_cell_text(cells[1], left_value, size_pt=9.0)

        if right is None:
            merged = cells[2].merge(cells[3])
            _set_cell_text(merged, "", color=_MUTED_COLOR, size_pt=9.0, allow_blank=True)
        else:
            _set_cell_shading(cells[2], _LABEL_FILL)
            _set_cell_text(cells[2], right_label, bold=True, color=_MUTED_COLOR, size_pt=8.8)
            _set_cell_text(cells[3], right_value, size_pt=9.0)

    def add_narrative_row(label: str, value: str) -> None:
        row = table.add_row()
        _prevent_row_split(row)
        cells = row.cells
        cells[0].width = Cm(2.8)
        _set_cell_shading(cells[0], _LABEL_FILL)
        _set_cell_text(cells[0], label, bold=True, color=_MUTED_COLOR, size_pt=8.8)
        merged = cells[1].merge(cells[3])
        merged.width = Cm(14.2)
        _set_cell_text(merged, value, size_pt=9.0)

    for label, value in fields:
        if _is_narrative_field(label, value):
            if compact_pending is not None:
                add_compact_row(compact_pending)
                compact_pending = None
            add_narrative_row(label, value)
            continue

        if compact_pending is None:
            compact_pending = (label, value)
        else:
            add_compact_row(compact_pending, (label, value))
            compact_pending = None

    if compact_pending is not None:
        add_compact_row(compact_pending)


def _table_width_weights(headers: Sequence[str], rows: Sequence[Sequence[str]]) -> list[float]:
    if not headers:
        return []
    weights: list[float] = []
    for index, header in enumerate(headers):
        lengths = [len(header or "-")]
        for row in rows[:20]:
            if index < len(row):
                lengths.append(len(row[index] or "-"))
        max_len = max(lengths)
        weights.append(float(min(26, max(7, max_len))))
    return weights


def _column_widths_cm(headers: Sequence[str], rows: Sequence[Sequence[str]], total_width_cm: float) -> list[float]:
    weights = _table_width_weights(headers, rows)
    if not weights:
        return []
    total_weight = sum(weights) or 1.0
    raw = [total_width_cm * weight / total_weight for weight in weights]
    # Avoid unusably narrow columns while preserving the total width approximately.
    minimum = 1.8 if len(headers) >= 7 else 2.2
    clamped = [max(minimum, width) for width in raw]
    scale = total_width_cm / sum(clamped)
    return [width * scale for width in clamped]


def _add_data_table(
    document: Any,
    table_block: Mapping[str, Any],
    *,
    page_width_cm: float = 17.0,
) -> None:
    caption = _text(table_block.get("caption"))
    headers = [_text(item) or "-" for item in _items(table_block.get("headers"))]
    rows = [[_text(item) or "-" for item in _items(row)] for row in _items(table_block.get("rows"))]

    if caption:
        paragraph = document.add_paragraph()
        _set_paragraph_spacing(paragraph, before=2, after=5, line=1.0)
        paragraph.paragraph_format.keep_with_next = True
        run = paragraph.add_run(caption)
        _set_run_style(run, size_pt=9.5, bold=True, color=_MUTED_COLOR)

    column_count = max(len(headers), max((len(row) for row in rows), default=0))
    if column_count <= 0:
        return

    table = document.add_table(rows=0, cols=column_count)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    _set_table_borders(table)

    normalized_headers = headers + ["-"] * (column_count - len(headers))
    normalized_rows = [row + ["-"] * (column_count - len(row)) for row in rows]
    widths = _column_widths_cm(normalized_headers, normalized_rows, page_width_cm)

    if normalized_headers:
        header_row = table.add_row()
        _set_repeat_table_header(header_row)
        _prevent_row_split(header_row)
        for index, value in enumerate(normalized_headers):
            cell = header_row.cells[index]
            if index < len(widths):
                cell.width = Cm(widths[index])
            _set_cell_shading(cell, _HEADER_FILL)
            _set_cell_text(cell, value, bold=True, color=_TEXT_COLOR, size_pt=8.4)

    for values in normalized_rows:
        row = table.add_row()
        _prevent_row_split(row)
        for index, value in enumerate(values):
            cell = row.cells[index]
            if index < len(widths):
                cell.width = Cm(widths[index])
            _set_cell_text(cell, value, size_pt=8.4)


def _section_needs_landscape(section: Mapping[str, Any]) -> bool:
    title = _text(section.get("title")).lower()
    for block_value in _items(section.get("blocks")):
        block = _mapping(block_value)
        if _text(block.get("type")) != "table":
            continue
        headers = _items(block.get("headers"))
        caption = _text(block.get("caption")).lower()
        header_text_length = sum(len(_text(item)) for item in headers)
        is_case_matrix = "case matrix" in caption or "case matrix" in title
        if len(headers) >= 7:
            return True
        if is_case_matrix and (len(headers) >= 6 or header_text_length >= 48):
            return True
    return False


def _switch_to_landscape(document: Any) -> None:
    section = document.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.8)

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_paragraph_spacing(paragraph)
    paragraph.clear()
    run = paragraph.add_run("해석 의뢰서  ·  ")
    _set_run_style(run, size_pt=8.0, color=_MUTED_COLOR)
    fld_simple = OxmlElement("w:fldSimple")
    fld_simple.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld_simple)


def build_word_docx(preview: Mapping[str, Any]) -> bytes:
    """Create a styled DOCX from the serialized current preview DOM, in DOM order.

    Input contract intentionally matches the existing browser serializer:
        {"sections": [{"title": str, "blocks": [...]}, ...]}

    Supported block types remain: field, stale, table.
    """
    _require_python_docx()

    document = Document()
    _configure_document(document)
    _add_document_title(document)

    landscape_started = False
    section_index = 0

    for section_value in _items(preview.get("sections")):
        section = _mapping(section_value)
        title = _text(section.get("title"))
        blocks = [_mapping(block) for block in _items(section.get("blocks"))]
        if not title and not blocks:
            continue

        if not landscape_started and _section_needs_landscape(section):
            _switch_to_landscape(document)
            landscape_started = True

        section_index += 1
        _add_section_heading(document, title or "구분", section_index)

        field_buffer: list[tuple[str, str]] = []

        def flush_fields() -> None:
            nonlocal field_buffer
            if field_buffer:
                _add_field_grid(document, field_buffer)
                field_buffer = []

        for block in blocks:
            kind = _text(block.get("type"))
            if kind == "field":
                field_buffer.append((_text(block.get("label")) or "항목", _text(block.get("value")) or "-"))
                continue

            flush_fields()

            if kind == "stale":
                _add_stale_notice(document, _text(block.get("text")) or "-")
            elif kind == "table":
                usable_width_cm = 26.1 if landscape_started else 17.0
                _add_data_table(document, block, page_width_cm=usable_width_cm)

        flush_fields()

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
