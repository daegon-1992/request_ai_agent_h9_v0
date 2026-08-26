#!/usr/bin/env python3
"""Build the structural-analysis SME request Word and Excel package.

The runtime application does not depend on this tool. It only creates the
handoff documents used before structural-analysis fields are implemented.
"""

from __future__ import annotations

import argparse
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, BadZipFile, ZipFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
OUTPUT_DIR = ROOT / "docs" / "structural_analysis_sme_request"
DOCX_PATH = OUTPUT_DIR / "구조해석_담당자_정보요청서.docx"
XLSX_PATH = OUTPUT_DIR / "구조해석_담당자_회신양식.xlsx"

PACKAGE_VERSION = "1.0"
EXPECTED_SHEETS = (
    "작성안내",
    "01_해석유형_우선순위",
    "02_완료사례_목록",
    "03_Field_Rule_Master",
    "04_핵심객체_규칙",
    "05_Case_조합규칙",
    "06_용어_안내_예외",
    "07_검수_승인",
    "08_변경이력",
    "09_워크숍_기록",
)

NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NS_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_PACKAGE_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"

ET.register_namespace("", NS_MAIN)
ET.register_namespace("r", NS_REL)


@dataclass(frozen=True)
class Validation:
    column: int
    values: tuple[str, ...]


@dataclass(frozen=True)
class SheetSpec:
    name: str
    headers: tuple[str, ...]
    example: tuple[str, ...] | None = None
    default_first_value: str = "실제 입력"
    blank_rows: int = 12
    validations: tuple[Validation, ...] = field(default_factory=tuple)
    rows: tuple[tuple[str, ...], ...] = field(default_factory=tuple)
    instruction_sheet: bool = False
    reference_sheet: bool = False


YES_NO = ("예", "아니오", "해당 없음", "검토 필요")
REVIEW_STATUS = ("작성 중", "검토 필요", "담당자 검수 완료", "최종 승인")


def _sheet_specs() -> tuple[SheetSpec, ...]:
    return (
        SheetSpec(
            name="작성안내",
            headers=("항목", "내용"),
            rows=(
                ("문서 목적", "비전문 사용자의 구조해석 의뢰를 담당자가 별도 재문의 없이 검토·착수할 수 있도록 업무 규칙을 수집합니다."),
                ("작성 순서", "먼저 01_해석유형_우선순위와 02_완료사례_목록을 작성합니다. 나머지 Master는 1차 워크숍 후 개발자가 초안을 작성하고 담당자가 검수합니다."),
                ("MVP 범위", "수행 빈도와 표준화 가능성이 높고 입력 확보가 쉬운 대표 해석유형 2~3종을 선정합니다."),
                ("작성 예시", "각 시트의 파란색 작성 예시 행은 사용법 안내이며 승인된 구조해석 규칙이 아닙니다."),
                ("입력 영역", "노란색 행에 실제 내용을 작성합니다. 확정할 수 없는 내용은 검토 필요로 표시합니다."),
                ("담당자 역할", "구조해석 업무 용어, 필수 여부, Case와 판정 규칙을 검수합니다. 영문 키·JSON·Python 구조는 개발자가 설계합니다."),
                ("의뢰자 필수 제외", "솔버, 메시, 접촉 알고리즘, 수렴 설정 등 담당자가 표준 절차로 결정할 수 있는 계산 제어값은 사용자 필수 입력에서 제외합니다."),
                ("사례 수", "MVP 후보 유형마다 정상, 다중 Case, 입력 누락·재문의 사례를 포함하여 최소 3건을 등록합니다."),
                ("민감정보", "개인정보, 프로젝트명, 도면번호 등은 익명화하거나 동일한 구조의 예시 데이터로 치환합니다."),
                ("승인 조건", "유형별 3개 사례를 Master로 재현하고 담당자가 무재문의 착수 가능 여부를 확인한 뒤 최종 승인합니다."),
                ("패키지 버전", PACKAGE_VERSION),
            ),
            instruction_sheet=True,
            blank_rows=0,
        ),
        SheetSpec(
            name="01_해석유형_우선순위",
            headers=(
                "행 구분", "사용자 표시명", "주요 질문", "적용 제품/부품", "대표 결과",
                "적용하지 않는 경우", "수행 빈도", "표준화 가능성", "입력 확보 난이도",
                "MVP 추천 여부", "검수 상태", "담당자 의견",
            ),
            example=(
                "작성 예시", "예: 구조해석 유형명", "제품이 요구 하중에서 안전한가?", "예: 제품군/어셈블리/부품",
                "예: 응력, 변위, 안전율", "예: 시간 이력이나 큰 변형이 핵심인 경우", "높음", "높음", "보통",
                "후보", "검토 필요", "예시는 실제 후보와 승인값으로 교체해 주세요.",
            ),
            validations=(
                Validation(1, ("작성 예시", "실제 입력")), Validation(7, ("높음", "중간", "낮음")),
                Validation(8, ("높음", "중간", "낮음")), Validation(9, ("쉬움", "보통", "어려움")),
                Validation(10, ("필수", "후보", "향후")), Validation(11, REVIEW_STATUS),
            ),
        ),
        SheetSpec(
            name="02_완료사례_목록",
            headers=(
                "행 구분", "사례 ID", "해석유형", "사례 구분", "최초 사용자 요청", "적용 제품/부품",
                "제품·CAD·도면 정보", "재질 정보", "하중 정보", "구속 정보", "접촉·체결 정보",
                "담당자 추가 질문", "최종 확정 조건", "최종 Case 목록", "요청 결과", "판정 기준",
                "최종 의뢰서·보고서 위치", "익명화 상태", "검수 상태", "담당자 의견",
            ),
            example=(
                "작성 예시", "CASE-예시-01", "예: 후보 유형명", "단일 하중·단일 형상", "최초 요청 문장을 가능한 그대로 기록",
                "예: 제품군/부품", "도면번호는 익명값 사용, CAD Revision과 단위 기록", "부품별 재질과 출처",
                "크기·단위·방향·적용 위치", "구속 위치와 자유도", "접촉·체결 방식", "실제 재문의 내용",
                "재문의 후 확정된 전체 조건", "Case 1: 기준 형상+하중 세트 A", "확인할 결과", "허용값·안전율·비교 기준",
                "익명화 자료 경로 또는 문서명", "대체 데이터", "검토 필요", "작성 예시이며 실제 사례가 아닙니다.",
            ),
            blank_rows=15,
            validations=(
                Validation(1, ("작성 예시", "실제 입력")),
                Validation(4, ("단일 하중·단일 형상", "다중 Case", "입력 누락·재문의", "기타")),
                Validation(18, ("완료", "대체 데이터", "검토 필요")), Validation(19, REVIEW_STATUS),
            ),
        ),
        SheetSpec(
            name="03_Field_Rule_Master",
            headers=(
                "행 구분", "규칙 ID(개발자 작성)", "입력 그룹", "사용자 표시명", "정의", "적용 해석유형",
                "필수 수준", "활성 조건", "입력 주체", "입력 방식", "단위·좌표계", "선택지·범위",
                "반복 가능 여부", "연결 대상", "모름/해당 없음 처리", "차단 검증", "사용자 질문",
                "정상 예시", "오류 예시", "검수 상태", "담당자 의견",
            ),
            example=(
                "작성 예시", "개발자 부여", "하중", "하중 크기", "제품에 작용하는 힘 또는 압력의 크기",
                "예: 후보 유형명", "조건부 필수", "하중 종류가 힘 또는 압력일 때", "의뢰자", "숫자",
                "N 또는 Pa, 기준 좌표계 명시", "0보다 큰 값, 직접 입력", "예", "하중 세트와 적용 면",
                "모름 허용, 담당자 협의 필요로 전환", "활성 상태에서 값·단위 누락 시 제출 차단",
                "어느 위치에 어떤 방향으로 얼마의 하중이 작용하나요?", "250 N, 제품 전면에서 +X 방향",
                "250처럼 단위와 방향이 없는 값", "검토 필요", "작성 예시는 승인 규칙이 아닙니다.",
            ),
            blank_rows=35,
            validations=(
                Validation(1, ("작성 예시", "실제 입력")),
                Validation(3, ("제품/형상", "재질", "연결", "구속", "하중", "환경", "결과 요청", "공통")),
                Validation(7, ("필수", "조건부 필수", "선택", "담당자 판단")),
                Validation(9, ("의뢰자", "구조 담당자", "시스템 계산")),
                Validation(10, ("숫자", "텍스트", "단일 선택", "다중 선택", "반복 행", "파일·이미지 표시", "날짜", "예/아니오")),
                Validation(13, YES_NO), Validation(20, REVIEW_STATUS),
            ),
        ),
        SheetSpec(
            name="04_핵심객체_규칙",
            headers=(
                "행 구분", "객체 규칙 ID(개발자 작성)", "객체 유형", "업무상 정의", "적용 해석유형",
                "최소 개수", "최대 개수", "복제 가능 여부", "구성 필드", "연결 대상", "금지 조합",
                "모름/해당 없음 처리", "입력 주체", "정상 예시", "검수 상태", "담당자 의견",
            ),
            example=(
                "작성 예시", "개발자 부여", "하중조건", "동시에 적용되는 하중의 묶음", "예: 후보 유형명",
                "1", "담당자 확인", "예", "종류, 크기, 단위, 방향, 적용 위치", "제품/부품/면과 Case",
                "예: 같은 자유도에 상충하는 강제변위", "위치를 모르면 이미지 표시를 요청", "의뢰자",
                "하중 세트 A: 중력+전면 250 N", "검토 필요", "작성 예시는 승인 규칙이 아닙니다.",
            ),
            blank_rows=24,
            validations=(
                Validation(1, ("작성 예시", "실제 입력")),
                Validation(3, ("제품/형상", "재질 할당", "연결/접촉", "구속조건", "하중조건", "환경조건", "결과 요청")),
                Validation(8, YES_NO), Validation(13, ("의뢰자", "구조 담당자", "시스템 계산", "공동 확인")),
                Validation(15, REVIEW_STATUS),
            ),
        ),
        SheetSpec(
            name="05_Case_조합규칙",
            headers=(
                "행 구분", "Case 규칙 ID(개발자 작성)", "적용 해석유형", "Case 구분 축", "동시 하중 여부",
                "별도 Case 분리 조건", "기본 Case 자동 생성", "금지 조합", "중복 Case 판정 기준",
                "필수 기준 Case", "Case 이름 입력 주체", "하중 조합계수 입력 주체", "안전계수 입력 주체",
                "정상 Case 예시", "검수 상태", "담당자 의견",
            ),
            example=(
                "작성 예시", "개발자 부여", "예: 후보 유형명", "형상, 하중 세트, 구속 세트", "조건에 따라 다름",
                "동시에 발생하지 않는 운전조건은 별도 Case", "검토 필요", "상충하는 구속과 강제변위",
                "동일 형상·재질·하중·구속·환경 조합", "기준 형상+기준 하중", "시스템 생성",
                "구조 담당자", "구조 담당자", "Case 1: 기준 형상/하중 A/구속 A", "검토 필요",
                "작성 예시는 승인 규칙이 아닙니다.",
            ),
            blank_rows=18,
            validations=(
                Validation(1, ("작성 예시", "실제 입력")),
                Validation(5, ("동시 적용", "별도 Case", "조건에 따라 다름", "검토 필요")),
                Validation(7, YES_NO),
                Validation(11, ("의뢰자", "구조 담당자", "시스템 생성", "공동 확인")),
                Validation(12, ("의뢰자", "구조 담당자", "시스템 계산", "해당 없음")),
                Validation(13, ("의뢰자", "구조 담당자", "시스템 계산", "해당 없음")),
                Validation(15, REVIEW_STATUS),
            ),
        ),
        SheetSpec(
            name="06_용어_안내_예외",
            headers=(
                "행 구분", "용어/필드", "쉬운 설명", "동의어·사용자 표현", "왜 필요한가", "기본 사용자 질문",
                "대체 질문", "Agent 추정 허용", "사용자 확인 필요", "모름/해당 없음 처리", "오류 수준",
                "잘못 선택 시 추천 유형", "도움말·정상 예시", "검수 상태", "담당자 의견",
            ),
            example=(
                "작성 예시", "구속조건", "제품이 움직이지 않도록 잡히거나 고정되는 위치와 방향",
                "고정부, 체결부, 지지부", "고정 방식에 따라 변형과 응력이 달라집니다.",
                "제품이 실제로 고정되거나 지지되는 위치는 어디인가요?", "도면이나 이미지에 고정 위치를 표시해 주세요.",
                "아니오", "예", "모름이면 담당자 협의 필요로 전환", "제출 차단", "해당 없음",
                "예: 바닥 체결부의 X/Y/Z 이동 고정", "검토 필요", "작성 예시는 승인 규칙이 아닙니다.",
            ),
            blank_rows=30,
            validations=(
                Validation(1, ("작성 예시", "실제 입력")), Validation(8, YES_NO), Validation(9, YES_NO),
                Validation(11, ("제출 차단", "경고", "안내", "해당 없음")), Validation(14, REVIEW_STATUS),
            ),
        ),
        SheetSpec(
            name="07_검수_승인",
            headers=(
                "행 구분", "해석유형", "활성 입력 그룹 확정", "필수·조건부 규칙 확정", "객체 연결 규칙 확정",
                "Case 규칙 확정", "모름·예외 처리 확정", "정상 사례 재현", "다중 Case 사례 재현",
                "입력 누락 사례 재현", "무재문의 착수 판단 가능", "구조 담당자", "최종 승인자",
                "승인일", "검수 상태", "미결 사항·의견",
            ),
            example=(
                "작성 예시", "예: MVP 유형명", "예", "예", "예", "예", "예", "예", "예", "예",
                "검토 필요", "담당자명", "승인자명", "YYYY-MM-DD", "검토 필요", "모든 항목이 예일 때만 최종 승인",
            ),
            blank_rows=10,
            validations=(
                Validation(1, ("작성 예시", "실제 입력")),
                *(Validation(column, YES_NO) for column in range(3, 12)),
                Validation(15, REVIEW_STATUS),
            ),
        ),
        SheetSpec(
            name="08_변경이력",
            headers=(
                "행 구분", "Master 버전", "변경일", "작성자", "변경 이유", "영향 시트", "영향 해석유형",
                "기존 의뢰서 호환 여부", "검토자", "검수 상태", "비고",
            ),
            example=(
                "작성 예시", "1.0", "YYYY-MM-DD", "작성자명", "최초 구조해석 Master 승인", "전체",
                "MVP 유형명", "해당 없음", "검토자명", "검토 필요", "승인 후 실제 변경이력을 기록",
            ),
            blank_rows=15,
            validations=(
                Validation(1, ("작성 예시", "실제 입력")), Validation(8, YES_NO), Validation(10, REVIEW_STATUS),
            ),
        ),
        SheetSpec(
            name="09_워크숍_기록",
            headers=(
                "행 구분", "회차", "일시", "참석자", "목표", "확정 결정", "미결 사항", "후속 담당자",
                "완료 예정일", "진행 상태", "비고",
            ),
            example=(
                "작성 예시", "1차", "YYYY-MM-DD HH:MM", "참석자명", "MVP 유형 선정과 대표 사례 분해",
                "확정된 내용만 기록", "추가 확인이 필요한 내용", "담당자명", "YYYY-MM-DD", "예정", "구두 미확정 내용은 결정으로 기록하지 않음",
            ),
            blank_rows=15,
            validations=(
                Validation(1, ("작성 예시", "실제 입력")),
                Validation(10, ("예정", "진행 중", "완료", "보류")),
            ),
        ),
    )


def _column_letter(number: int) -> str:
    letters = ""
    while number:
        number, remainder = divmod(number - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


def _xml_bytes(element: ET.Element) -> bytes:
    return ET.tostring(element, encoding="utf-8", xml_declaration=True)


def _inline_cell(row: ET.Element, reference: str, value: str, style: int) -> None:
    cell = ET.SubElement(row, f"{{{NS_MAIN}}}c", {"r": reference, "t": "inlineStr", "s": str(style)})
    inline = ET.SubElement(cell, f"{{{NS_MAIN}}}is")
    text = ET.SubElement(inline, f"{{{NS_MAIN}}}t", {XML_SPACE: "preserve"})
    text.text = value


def _column_width(values: Iterable[str]) -> float:
    longest = max((max((len(line) for line in str(value).splitlines()), default=0) for value in values), default=8)
    return float(min(max(longest + 2, 11), 42))


def _worksheet_xml(spec: SheetSpec) -> bytes:
    worksheet = ET.Element(f"{{{NS_MAIN}}}worksheet")
    views = ET.SubElement(worksheet, f"{{{NS_MAIN}}}sheetViews")
    view = ET.SubElement(views, f"{{{NS_MAIN}}}sheetView", {"workbookViewId": "0"})
    if not spec.instruction_sheet:
        ET.SubElement(view, f"{{{NS_MAIN}}}pane", {
            "ySplit": "1", "topLeftCell": "A2", "activePane": "bottomLeft", "state": "frozen",
        })
    ET.SubElement(worksheet, f"{{{NS_MAIN}}}sheetFormatPr", {"defaultRowHeight": "20"})

    if spec.instruction_sheet:
        data_rows = list(spec.rows)
    else:
        data_rows = []
        if spec.example:
            data_rows.append(spec.example)
        data_rows.extend(spec.rows)
        for _ in range(spec.blank_rows):
            data_rows.append((spec.default_first_value, *("" for _ in spec.headers[1:])))

    columns = ET.SubElement(worksheet, f"{{{NS_MAIN}}}cols")
    for index, header in enumerate(spec.headers, 1):
        values = [header, *(row[index - 1] if index - 1 < len(row) else "" for row in data_rows)]
        width = 24.0 if spec.instruction_sheet and index == 2 else _column_width(values)
        ET.SubElement(columns, f"{{{NS_MAIN}}}col", {
            "min": str(index), "max": str(index), "width": str(width), "customWidth": "1",
        })

    sheet_data = ET.SubElement(worksheet, f"{{{NS_MAIN}}}sheetData")
    header_row = ET.SubElement(sheet_data, f"{{{NS_MAIN}}}row", {"r": "1", "ht": "34", "customHeight": "1"})
    for index, header in enumerate(spec.headers, 1):
        _inline_cell(header_row, f"{_column_letter(index)}1", header, 1)

    for row_index, values in enumerate(data_rows, 2):
        if spec.instruction_sheet:
            row_style = 4 if values[0] in {"문서 목적", "작성 순서", "MVP 범위", "승인 조건"} else 6
        else:
            if row_index == 2 and spec.example:
                row_style = 3
            elif spec.reference_sheet and row_index <= 1 + bool(spec.example) + len(spec.rows):
                row_style = 6
            else:
                row_style = 2
        height = "48" if row_style == 3 else "36"
        row = ET.SubElement(sheet_data, f"{{{NS_MAIN}}}row", {"r": str(row_index), "ht": height, "customHeight": "1"})
        for column_index in range(1, len(spec.headers) + 1):
            value = values[column_index - 1] if column_index - 1 < len(values) else ""
            _inline_cell(row, f"{_column_letter(column_index)}{row_index}", value, row_style)

    last_row = max(1, len(data_rows) + 1)
    last_column = _column_letter(len(spec.headers))
    ET.SubElement(worksheet, f"{{{NS_MAIN}}}autoFilter", {"ref": f"A1:{last_column}{last_row}"})

    if spec.validations:
        validations = ET.SubElement(worksheet, f"{{{NS_MAIN}}}dataValidations", {"count": str(len(spec.validations))})
        for validation in spec.validations:
            column = _column_letter(validation.column)
            data_validation = ET.SubElement(validations, f"{{{NS_MAIN}}}dataValidation", {
                "type": "list",
                "allowBlank": "1",
                "showErrorMessage": "1",
                "errorStyle": "stop",
                "errorTitle": "허용되지 않는 값",
                "error": "목록에서 값을 선택해 주세요.",
                "sqref": f"{column}2:{column}{last_row}",
            })
            formula = ET.SubElement(data_validation, f"{{{NS_MAIN}}}formula1")
            formula.text = f'"{",".join(validation.values)}"'

    ET.SubElement(worksheet, f"{{{NS_MAIN}}}pageMargins", {
        "left": "0.3", "right": "0.3", "top": "0.5", "bottom": "0.5", "header": "0.2", "footer": "0.2",
    })
    return _xml_bytes(worksheet)


def _styles_xml() -> bytes:
    root = ET.Element(f"{{{NS_MAIN}}}styleSheet")
    fonts = ET.SubElement(root, f"{{{NS_MAIN}}}fonts", {"count": "4"})
    default_font = ET.SubElement(fonts, f"{{{NS_MAIN}}}font")
    ET.SubElement(default_font, f"{{{NS_MAIN}}}sz", {"val": "10"})
    ET.SubElement(default_font, f"{{{NS_MAIN}}}name", {"val": "맑은 고딕"})
    header_font = ET.SubElement(fonts, f"{{{NS_MAIN}}}font")
    ET.SubElement(header_font, f"{{{NS_MAIN}}}b")
    ET.SubElement(header_font, f"{{{NS_MAIN}}}color", {"rgb": "FFFFFFFF"})
    ET.SubElement(header_font, f"{{{NS_MAIN}}}sz", {"val": "10"})
    ET.SubElement(header_font, f"{{{NS_MAIN}}}name", {"val": "맑은 고딕"})
    bold_font = ET.SubElement(fonts, f"{{{NS_MAIN}}}font")
    ET.SubElement(bold_font, f"{{{NS_MAIN}}}b")
    ET.SubElement(bold_font, f"{{{NS_MAIN}}}color", {"rgb": "FF17324D"})
    ET.SubElement(bold_font, f"{{{NS_MAIN}}}sz", {"val": "10"})
    ET.SubElement(bold_font, f"{{{NS_MAIN}}}name", {"val": "맑은 고딕"})
    note_font = ET.SubElement(fonts, f"{{{NS_MAIN}}}font")
    ET.SubElement(note_font, f"{{{NS_MAIN}}}i")
    ET.SubElement(note_font, f"{{{NS_MAIN}}}color", {"rgb": "FF5F6B76"})
    ET.SubElement(note_font, f"{{{NS_MAIN}}}sz", {"val": "10"})
    ET.SubElement(note_font, f"{{{NS_MAIN}}}name", {"val": "맑은 고딕"})

    fills = ET.SubElement(root, f"{{{NS_MAIN}}}fills", {"count": "6"})
    ET.SubElement(ET.SubElement(fills, f"{{{NS_MAIN}}}fill"), f"{{{NS_MAIN}}}patternFill", {"patternType": "none"})
    ET.SubElement(ET.SubElement(fills, f"{{{NS_MAIN}}}fill"), f"{{{NS_MAIN}}}patternFill", {"patternType": "gray125"})
    for color in ("FF17324D", "FFFFF2CC", "FFDDEBF7", "FFE2F0D9"):
        fill = ET.SubElement(fills, f"{{{NS_MAIN}}}fill")
        pattern = ET.SubElement(fill, f"{{{NS_MAIN}}}patternFill", {"patternType": "solid"})
        ET.SubElement(pattern, f"{{{NS_MAIN}}}fgColor", {"rgb": color})
        ET.SubElement(pattern, f"{{{NS_MAIN}}}bgColor", {"indexed": "64"})

    borders = ET.SubElement(root, f"{{{NS_MAIN}}}borders", {"count": "2"})
    ET.SubElement(borders, f"{{{NS_MAIN}}}border")
    border = ET.SubElement(borders, f"{{{NS_MAIN}}}border")
    for edge in ("left", "right", "top", "bottom"):
        side = ET.SubElement(border, f"{{{NS_MAIN}}}{edge}", {"style": "thin"})
        ET.SubElement(side, f"{{{NS_MAIN}}}color", {"rgb": "FFD9E1E8"})
    ET.SubElement(border, f"{{{NS_MAIN}}}diagonal")

    ET.SubElement(root, f"{{{NS_MAIN}}}cellStyleXfs", {"count": "1"}).append(
        ET.Element(f"{{{NS_MAIN}}}xf", {"numFmtId": "0", "fontId": "0", "fillId": "0", "borderId": "0"})
    )
    cell_xfs = ET.SubElement(root, f"{{{NS_MAIN}}}cellXfs", {"count": "7"})
    ET.SubElement(cell_xfs, f"{{{NS_MAIN}}}xf", {"numFmtId": "0", "fontId": "0", "fillId": "0", "borderId": "0", "xfId": "0"})
    style_defs = (
        ("1", "2", "1", "center"),
        ("0", "3", "1", "left"),
        ("3", "4", "1", "left"),
        ("2", "5", "1", "left"),
        ("3", "0", "0", "left"),
        ("0", "0", "1", "left"),
    )
    for font_id, fill_id, border_id, horizontal in style_defs:
        xf = ET.SubElement(cell_xfs, f"{{{NS_MAIN}}}xf", {
            "numFmtId": "0", "fontId": font_id, "fillId": fill_id, "borderId": border_id,
            "xfId": "0", "applyAlignment": "1", "applyFill": "1", "applyBorder": "1",
        })
        ET.SubElement(xf, f"{{{NS_MAIN}}}alignment", {
            "horizontal": horizontal, "vertical": "top", "wrapText": "1",
        })

    cell_styles = ET.SubElement(root, f"{{{NS_MAIN}}}cellStyles", {"count": "1"})
    ET.SubElement(cell_styles, f"{{{NS_MAIN}}}cellStyle", {"name": "Normal", "xfId": "0", "builtinId": "0"})
    ET.SubElement(root, f"{{{NS_MAIN}}}dxfs", {"count": "0"})
    ET.SubElement(root, f"{{{NS_MAIN}}}tableStyles", {
        "count": "0", "defaultTableStyle": "TableStyleMedium2", "defaultPivotStyle": "PivotStyleLight16",
    })
    return _xml_bytes(root)


def _workbook_xml(specs: tuple[SheetSpec, ...]) -> bytes:
    root = ET.Element(f"{{{NS_MAIN}}}workbook")
    ET.SubElement(root, f"{{{NS_MAIN}}}bookViews").append(
        ET.Element(f"{{{NS_MAIN}}}workbookView", {"xWindow": "0", "yWindow": "0", "windowWidth": "24000", "windowHeight": "12000"})
    )
    sheets = ET.SubElement(root, f"{{{NS_MAIN}}}sheets")
    for index, spec in enumerate(specs, 1):
        ET.SubElement(sheets, f"{{{NS_MAIN}}}sheet", {
            "name": spec.name, "sheetId": str(index), f"{{{NS_REL}}}id": f"rId{index}",
        })
    ET.SubElement(root, f"{{{NS_MAIN}}}calcPr", {"calcId": "191029", "fullCalcOnLoad": "1"})
    return _xml_bytes(root)


def _workbook_rels(specs: tuple[SheetSpec, ...]) -> bytes:
    root = ET.Element("Relationships", {"xmlns": NS_PACKAGE_REL})
    for index in range(1, len(specs) + 1):
        ET.SubElement(root, "Relationship", {
            "Id": f"rId{index}",
            "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet",
            "Target": f"worksheets/sheet{index}.xml",
        })
    ET.SubElement(root, "Relationship", {
        "Id": f"rId{len(specs) + 1}",
        "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles",
        "Target": "styles.xml",
    })
    return _xml_bytes(root)


def _content_types(specs: tuple[SheetSpec, ...]) -> bytes:
    root = ET.Element("Types", {"xmlns": "http://schemas.openxmlformats.org/package/2006/content-types"})
    ET.SubElement(root, "Default", {"Extension": "rels", "ContentType": "application/vnd.openxmlformats-package.relationships+xml"})
    ET.SubElement(root, "Default", {"Extension": "xml", "ContentType": "application/xml"})
    ET.SubElement(root, "Override", {"PartName": "/xl/workbook.xml", "ContentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"})
    ET.SubElement(root, "Override", {"PartName": "/xl/styles.xml", "ContentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"})
    for index in range(1, len(specs) + 1):
        ET.SubElement(root, "Override", {
            "PartName": f"/xl/worksheets/sheet{index}.xml",
            "ContentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml",
        })
    ET.SubElement(root, "Override", {"PartName": "/docProps/core.xml", "ContentType": "application/vnd.openxmlformats-package.core-properties+xml"})
    ET.SubElement(root, "Override", {"PartName": "/docProps/app.xml", "ContentType": "application/vnd.openxmlformats-officedocument.extended-properties+xml"})
    return _xml_bytes(root)


def _root_rels() -> bytes:
    root = ET.Element("Relationships", {"xmlns": NS_PACKAGE_REL})
    relationships = (
        ("rId1", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument", "xl/workbook.xml"),
        ("rId2", "http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties", "docProps/core.xml"),
        ("rId3", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/extended-properties", "docProps/app.xml"),
    )
    for rel_id, rel_type, target in relationships:
        ET.SubElement(root, "Relationship", {"Id": rel_id, "Type": rel_type, "Target": target})
    return _xml_bytes(root)


def _core_properties() -> bytes:
    xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/" xmlns:dcmitype="http://purl.org/dc/dcmitype/" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
  <dc:title>구조해석 담당자 회신양식</dc:title>
  <dc:subject>h7_v0 구조해석 확장 사전 정보수집</dc:subject>
  <dc:creator>CAE Request Agent</dc:creator>
  <cp:lastModifiedBy>CAE Request Agent</cp:lastModifiedBy>
</cp:coreProperties>"""
    return xml.encode("utf-8")


def _app_properties(specs: tuple[SheetSpec, ...]) -> bytes:
    titles = "".join(f"<vt:lpstr>{spec.name}</vt:lpstr>" for spec in specs)
    xml = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">
  <Application>CAE Request Agent</Application>
  <DocSecurity>0</DocSecurity>
  <ScaleCrop>false</ScaleCrop>
  <HeadingPairs><vt:vector size="2" baseType="variant"><vt:variant><vt:lpstr>Worksheets</vt:lpstr></vt:variant><vt:variant><vt:i4>{len(specs)}</vt:i4></vt:variant></vt:vector></HeadingPairs>
  <TitlesOfParts><vt:vector size="{len(specs)}" baseType="lpstr">{titles}</vt:vector></TitlesOfParts>
  <Company>CAE Request Agent</Company>
  <AppVersion>1.0</AppVersion>
</Properties>"""
    return xml.encode("utf-8")


def write_xlsx(path: Path, specs: tuple[SheetSpec, ...]) -> None:
    """Write a styled workbook from the shared request-pack sheet contract."""

    path.parent.mkdir(parents=True, exist_ok=True)
    with ZipFile(path, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", _content_types(specs))
        archive.writestr("_rels/.rels", _root_rels())
        archive.writestr("docProps/core.xml", _core_properties())
        archive.writestr("docProps/app.xml", _app_properties(specs))
        archive.writestr("xl/workbook.xml", _workbook_xml(specs))
        archive.writestr("xl/_rels/workbook.xml.rels", _workbook_rels(specs))
        archive.writestr("xl/styles.xml", _styles_xml())
        for index, spec in enumerate(specs, 1):
            archive.writestr(f"xl/worksheets/sheet{index}.xml", _worksheet_xml(spec))


def build_xlsx() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    write_xlsx(XLSX_PATH, _sheet_specs())


def _set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_cell_text(cell, text: str, *, bold: bool = False, color: str = "222222") -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "맑은 고딕"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _add_table(document: Document, headers: tuple[str, ...], rows: tuple[tuple[str, ...], ...], widths: tuple[float, ...] | None = None) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        _set_cell_text(table.rows[0].cells[index], header, bold=True, color="FFFFFF")
        _set_cell_shading(table.rows[0].cells[index], "17324D")
    for row_values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            _set_cell_text(cells[index], value)
            if index == 0:
                _set_cell_shading(cells[index], "E2F0D9")
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Cm(width)
    document.add_paragraph()


def _add_bullets(document: Document, items: Iterable[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


def build_docx() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    normal = document.styles["Normal"]
    normal.font.name = "맑은 고딕"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.2
    for style_name in ("Title", "Heading 1", "Heading 2"):
        style = document.styles[style_name]
        style.font.name = "맑은 고딕"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
        style.font.color.rgb = RGBColor(23, 50, 77)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("구조해석 의뢰 Agent 확장을 위한\n자료 및 검토 요청서")
    run.bold = True
    run.font.name = "맑은 고딕"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(23, 50, 77)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(f"h7_v0 구조해석 확장 사전 정보수집 · 패키지 버전 {PACKAGE_VERSION}").italic = True
    document.add_paragraph()

    _add_table(document, ("항목", "기입 내용"), (
        ("수신", "[구조해석 담당자/조직 입력]"),
        ("발신", "[요청자/조직 입력]"),
        ("회신 희망일", "[YYYY-MM-DD]"),
        ("1차 워크숍", "[후보 일정 입력]"),
    ), (3.2, 12.8))

    document.add_heading("1. 요청 배경과 목표", level=1)
    document.add_paragraph(
        "비전문 사용자도 구조해석을 쉽게 의뢰할 수 있도록 기존 유동해석 의뢰 Agent를 확장하려고 합니다. "
        "목표는 구조해석 계산 설정을 완전히 자동화하는 것이 아니라, 사용자가 목적·대상 제품·재질·하중·구속·연결 상태와 "
        "결과 판정기준을 빠짐없이 전달하여 구조해석 담당자가 별도 재문의 없이 해석 착수 여부를 판단할 수 있게 하는 것입니다."
    )
    document.add_paragraph(
        "초기에는 수행 빈도와 표준화 가능성이 높고 일반 의뢰자가 입력을 확보하기 쉬운 대표 해석유형 2~3종만 선정합니다. "
        "기존 사업부·제품군·Platform 및 Base/비교 제품 흐름은 재사용합니다."
    )

    document.add_heading("2. 요청드리는 자료", level=1)
    _add_bullets(document, (
        "후보 구조해석 유형과 업무상 우선순위",
        "후보 유형별 익명화된 실제 의뢰서·보고서·입력자료 3건 이상",
        "의뢰자가 반드시 제공해야 하는 정보와 구조 담당자가 판단해야 하는 정보의 구분",
        "재질·하중·구속·접촉/체결·해석 Case·결과 판정 규칙",
        "정보가 없거나 애매할 때 실제로 확인하는 후속 질문",
        "개발자가 작성한 Field Rule Master와 샘플 의뢰서의 최종 검수",
    ))

    document.add_heading("3. 우선 작성 범위", level=1)
    document.add_paragraph(
        "첨부된 Excel 회신양식에서 먼저 `01_해석유형_우선순위`와 `02_완료사례_목록` 시트를 작성해 주세요. "
        "나머지 Master 시트는 1차 워크숍에서 실제 사례 한 건을 함께 검토한 뒤 개발자가 초안을 작성합니다."
    )
    _add_table(document, ("사례 구분", "포함할 내용"), (
        ("정상 사례", "일반적인 단일 하중·단일 형상 사례"),
        ("다중 Case", "여러 하중, 형상 또는 조건을 비교한 사례"),
        ("정보 부족 사례", "입력 누락, 조건부 필드 또는 담당자 재문의가 있었던 사례"),
    ), (4.0, 12.0))

    document.add_heading("4. Master에서 확인할 업무 규칙", level=1)
    _add_table(document, ("구분", "확인 내용"), (
        ("제품/형상", "어셈블리·부품 식별자, CAD Revision, 기준/비교 형상, 단순화 부위, 관심 위치"),
        ("재질", "대상 부품, 재질명, 필요한 물성, 물성 출처와 담당자 보완 규칙"),
        ("연결/접촉", "연결 부품, 접촉·용접·볼트·체결 유형과 필수 파라미터"),
        ("구속", "적용 위치, 제한 자유도, 좌표계와 실제 설치·지지 이유"),
        ("하중", "종류, 크기, 단위, 방향, 작용 위치, 시간·주파수 의존성과 동시 적용 관계"),
        ("환경", "온도, 중력 방향, 조립·설치 상태 등 해석유형별 필요 항목"),
        ("결과 요청", "결과 종류, 확인 위치, 비교 대상, 허용 기준, 안전율 또는 합격 조건"),
        ("Case", "구분 축, 동시 하중, 기본 Case, 금지·중복 조합, 기준 Case와 계수 입력 주체"),
    ), (4.0, 12.0))

    document.add_heading("5. 협업 방식", level=1)
    _add_bullets(document, (
        "사전 회신: 후보 유형과 실제 사례를 먼저 공유합니다.",
        "1차 워크숍: MVP 2~3종을 선정하고 대표 사례 한 건을 처음부터 끝까지 분해합니다.",
        "개발자 초안: 사례를 Field Rule, 반복 객체, Case 규칙과 사용자 질문으로 변환합니다.",
        "2차 워크숍: 필수·조건부·담당자 판단 항목과 예외를 확정합니다.",
        "최종 검수: 유형별 3개 사례를 Master로 재현하고 무재문의 착수 가능 여부를 승인합니다.",
    ))

    document.add_heading("6. 구조 담당자에게 요청하지 않는 사항", level=1)
    document.add_paragraph(
        "구조 담당자께 영문 필드 키, JSON, Python 또는 데이터베이스 구조 설계를 요청드리지 않습니다. "
        "구조 담당자는 한글 업무 규칙을 확정하고 개발자가 시스템 구조로 변환합니다."
    )
    _add_bullets(document, (
        "솔버 및 해석 코드 선택",
        "요소 종류와 메시 크기",
        "접촉 알고리즘의 수치 설정",
        "비선형 증분과 수렴 설정",
        "담당자가 표준 절차로 결정할 수 있는 기타 계산 제어값",
    ))

    document.add_heading("7. 승인 기준", level=1)
    _add_bullets(document, (
        "MVP 유형마다 활성 입력 그룹과 필수·조건부 필드가 결정되어 있습니다.",
        "재질·하중·구속·연결·결과 요청과 제품·부품·Case의 연결이 명확합니다.",
        "사용자가 모르는 값을 입력했을 때 후속 질문 또는 담당자 이관 방식이 있습니다.",
        "정상·다중 Case·정보 부족 사례가 승인된 Master로 재현됩니다.",
        "구조 담당자가 추가 전화나 메일 없이 착수 여부를 판단할 수 있습니다.",
    ))

    document.add_heading("8. 회신 정보", level=1)
    _add_table(document, ("항목", "회신"), (
        ("구조해석 업무 담당자", ""),
        ("Master 최종 승인자", ""),
        ("1차 워크숍 가능 일정", ""),
        ("보안·익명화 유의사항", ""),
        ("기타 의견", ""),
    ), (5.0, 11.0))

    document.core_properties.title = "구조해석 담당자 정보요청서"
    document.core_properties.subject = "h7_v0 구조해석 확장 사전 정보수집"
    document.core_properties.author = "CAE Request Agent"
    document.save(DOCX_PATH)


def check_outputs() -> list[str]:
    errors: list[str] = []
    required_markdown = (
        "README.md",
        "01_구조해석_담당자_요청문안.md",
        "02_작성_워크숍_검수_가이드.md",
        "03_회신_데이터_계약.md",
    )
    for filename in required_markdown:
        path = OUTPUT_DIR / filename
        if not path.exists():
            errors.append(f"missing file: {path}")
            continue
        try:
            path.read_text(encoding="utf-8", errors="strict")
        except UnicodeDecodeError as exc:
            errors.append(f"invalid UTF-8: {path}: {exc}")

    try:
        workbook = Document(DOCX_PATH)
        docx_text = "\n".join(paragraph.text for paragraph in workbook.paragraphs)
        for phrase in ("요청 배경과 목표", "협업 방식", "승인 기준"):
            if phrase not in docx_text:
                errors.append(f"DOCX missing phrase: {phrase}")
    except (OSError, BadZipFile, ValueError) as exc:
        errors.append(f"invalid DOCX: {exc}")

    try:
        with ZipFile(XLSX_PATH) as archive:
            names = set(archive.namelist())
            required_entries = {"[Content_Types].xml", "xl/workbook.xml", "xl/styles.xml"}
            missing_entries = required_entries - names
            if missing_entries:
                errors.append(f"XLSX missing entries: {sorted(missing_entries)}")
            workbook_root = ET.fromstring(archive.read("xl/workbook.xml"))
            sheet_names = tuple(
                sheet.attrib["name"]
                for sheet in workbook_root.findall(f".//{{{NS_MAIN}}}sheet")
            )
            if sheet_names != EXPECTED_SHEETS:
                errors.append(f"unexpected XLSX sheets: {sheet_names}")
            for index in range(1, len(EXPECTED_SHEETS) + 1):
                ET.fromstring(archive.read(f"xl/worksheets/sheet{index}.xml"))
            ET.fromstring(archive.read("xl/styles.xml"))
    except (OSError, KeyError, BadZipFile, ET.ParseError) as exc:
        errors.append(f"invalid XLSX: {exc}")
    return errors


def main() -> int:
    parser = argparse.ArgumentParser(description="Build or verify the structural-analysis SME request package.")
    parser.add_argument("--check", action="store_true", help="verify existing outputs without rewriting them")
    args = parser.parse_args()

    if not args.check:
        build_docx()
        build_xlsx()
        print(f"Generated: {DOCX_PATH}")
        print(f"Generated: {XLSX_PATH}")

    errors = check_outputs()
    if errors:
        print("Package verification failed:")
        for error in errors:
            print(f"- {error}")
        return 1
    print("Package verification passed")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
